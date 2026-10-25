"""
exportar.py
===========
Exporta los resultados a archivos que el usuario puede descargar:

    generar_excel(res)       -> reporte .xlsx con formato y gráficos insertados
    generar_pdf(res)         -> reporte .pdf (portada, tablas y gráficos)
    generar_zip_imagenes(res)-> .zip con todos los gráficos en PNG

Todas devuelven BYTES (el archivo "en memoria"). Streamlit los entrega con
st.download_button(...) sin necesidad de guardar nada en el disco del servidor.
"""

import io
import zipfile
from datetime import datetime

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .calculos import ResultadosAnalisis
from .graficos import figura_a_png, generar_todas_las_figuras

TITULO = "ANÁLISIS SÍSMICO DINÁMICO MODAL ESPECTRAL (NTE E.030)"

# ---- Estilos de Excel --------------------------------------------------------
AZUL = "1B365D"
RELLENO_ENCABEZADO = PatternFill(start_color=AZUL, end_color=AZUL, fill_type="solid")
RELLENO_CEBRA = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
FUENTE_ENCABEZADO = Font(bold=True, color="FFFFFF")
FUENTE_TITULO = Font(size=15, bold=True, color=AZUL)
FUENTE_SUBTITULO = Font(size=10, italic=True, color="555555")
BORDE = Border(*(Side(style="thin", color="D3D3D3"),) * 4)


def _escribir_tabla(ws, df, fila_inicio: int, titulo: str = None) -> int:
    """Escribe un DataFrame con formato en la hoja `ws`. Devuelve la siguiente fila libre."""
    if titulo:
        ws.cell(row=fila_inicio, column=1, value=titulo).font = Font(bold=True, size=12, color=AZUL)
        fila_inicio += 1

    for c, nombre in enumerate(df.columns, start=1):
        celda = ws.cell(row=fila_inicio, column=c, value=str(nombre))
        celda.fill, celda.font, celda.border = RELLENO_ENCABEZADO, FUENTE_ENCABEZADO, BORDE
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r, fila in enumerate(df.itertuples(index=False), start=fila_inicio + 1):
        for c, valor in enumerate(fila, start=1):
            if hasattr(valor, "item"):        # numpy -> tipo nativo de Python
                valor = valor.item()
            celda = ws.cell(row=r, column=c, value=valor)
            celda.border = BORDE
            if isinstance(valor, float):
                celda.number_format = "0.0000" if abs(valor) < 1 else "#,##0.000"
            if (r - fila_inicio) % 2 == 0:
                celda.fill = RELLENO_CEBRA
    return fila_inicio + len(df) + 2


def _ajustar_anchos(ws, minimo: int = 12, maximo: int = 45):
    for columna in ws.columns:
        largo = max(len(str(c.value)) if c.value is not None else 0 for c in columna)
        ws.column_dimensions[get_column_letter(columna[0].column)].width = max(minimo, min(largo + 3, maximo))


def generar_excel(res: ResultadosAnalisis, figuras: dict = None) -> bytes:
    """Reporte Excel: Resumen, Resultados modales, Resultados por piso y Gráficos."""
    figuras = figuras or generar_todas_las_figuras(res)
    wb = Workbook()

    # --- Hoja 1: Resumen ---------------------------------------------------
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"], ws["A1"].font = TITULO, FUENTE_TITULO
    ws["A2"] = f"Edificación de {len(res.pisos)} pisos · {res.n_modos} modos · generado {datetime.now():%Y-%m-%d %H:%M}"
    ws["A2"].font = FUENTE_SUBTITULO
    fila = _escribir_tabla(ws, res.tabla_parametros(), 4, "Parámetros sísmicos")
    opciones = pd.DataFrame([
        ("Modos normalizados respecto a M", "Sí" if res.opciones.get("normalizados", True) else "No"),
        ("Peso Σ|V| en la combinación", res.opciones.get("peso_abs", 0.25)),
        ("Peso √ΣV² en la combinación", res.opciones.get("peso_rcsc", 0.75)),
        ("Factor de la cortante real", res.opciones.get("factor_vreal", 0.75)),
    ], columns=["Opción", "Valor"])
    fila = _escribir_tabla(ws, opciones, fila, "Opciones de cálculo")
    resumen = pd.DataFrame(list(res.resumen().items()), columns=["Resultado", "Valor"])
    _escribir_tabla(ws, resumen, fila, "Resultados clave")
    _ajustar_anchos(ws)

    # --- Hoja 2: Análisis modal --------------------------------------------
    ws = wb.create_sheet("Analisis modal")
    ws["A1"], ws["A1"].font = "PROPIEDADES MODALES Y ESPECTRALES", FUENTE_TITULO
    fila = _escribir_tabla(ws, res.tabla_modal(), 3, "Resumen modal")
    fila = _escribir_tabla(ws, res._tabla_por_piso(res.phi), fila, "Formas modales φ")
    fila = _escribir_tabla(ws, res.tabla_aceleraciones(), fila, "Aceleraciones por piso Ü (m/s²)")
    fila = _escribir_tabla(ws, res.tabla_fuerzas(), fila, "Fuerzas sísmicas por piso F (t)")
    _escribir_tabla(ws, res.tabla_cortantes_modales(), fila, "Cortantes por piso V (t)")
    _ajustar_anchos(ws)

    # --- Hoja 3: Cortantes combinados ---------------------------------------
    ws = wb.create_sheet("Resultados por piso")
    ws["A1"], ws["A1"].font = "CORTANTES SÍSMICAS POR PISO Y COMBINACIÓN MODAL", FUENTE_TITULO
    _escribir_tabla(ws, res.tabla_combinacion(), 3)
    _ajustar_anchos(ws)

    # --- Hoja 4: Gráficos -------------------------------------------------
    ws = wb.create_sheet("Graficos")
    fila = 1
    for nombre, fig in figuras.items():
        img = XLImage(io.BytesIO(figura_a_png(fig, dpi=110)))
        escala = 720 / img.width
        img.width, img.height = img.width * escala, img.height * escala
        ws.cell(row=fila, column=1, value=nombre).font = Font(bold=True, color=AZUL)
        ws.add_image(img, f"A{fila + 1}")
        fila += int(img.height / 20) + 4   # ~20 px por fila

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---- PDF -------------------------------------------------------------------
def _pagina_tabla(pdf, df, titulo, decimales=4):
    """Dibuja un DataFrame como tabla en una página A4 horizontal."""
    fig, ax = plt.subplots(figsize=(11.69, 8.27))
    ax.axis("off")
    ax.set_title(titulo, fontsize=14, fontweight="bold", color="#1B365D", loc="left")
    texto = df.copy()
    for col in texto.columns:
        if texto[col].dtype.kind == "f":
            texto[col] = texto[col].map(lambda v: f"{v:,.{decimales}f}")
    tabla = ax.table(cellText=texto.values, colLabels=list(texto.columns),
                     loc="upper center", cellLoc="center")
    tabla.auto_set_font_size(False)
    tabla.set_fontsize(8.5)
    tabla.scale(1, 1.6)
    for (r, _), celda in tabla.get_celld().items():
        if r == 0:
            celda.set_facecolor("#1B365D")
            celda.set_text_props(color="white", fontweight="bold")
        elif r % 2 == 0:
            celda.set_facecolor("#F2F5F9")
    pdf.savefig(fig)
    plt.close(fig)


def generar_pdf(res: ResultadosAnalisis, figuras: dict = None) -> bytes:
    """Reporte PDF: portada con parámetros, tablas de resultados y todos los gráficos."""
    figuras = figuras or generar_todas_las_figuras(res)
    buffer = io.BytesIO()
    with PdfPages(buffer) as pdf:
        # Portada
        fig, ax = plt.subplots(figsize=(11.69, 8.27))
        ax.axis("off")
        ax.text(0.0, 0.95, TITULO, fontsize=18, fontweight="bold", color="#1B365D", transform=ax.transAxes)
        ax.text(0.0, 0.89, f"Generado: {datetime.now():%Y-%m-%d %H:%M}  ·  {len(res.pisos)} pisos  ·  {res.n_modos} modos",
                fontsize=10, color="gray", transform=ax.transAxes)
        lineas = [f"{k} = {v:g}" for k, v in res.params.como_diccionario().items()]
        lineas += ["", *[f"{k}: {v:,.3f}" for k, v in res.resumen().items()]]
        ax.text(0.0, 0.80, "\n".join(lineas), fontsize=12, va="top", family="monospace",
                transform=ax.transAxes)
        pdf.savefig(fig)
        plt.close(fig)

        _pagina_tabla(pdf, res.tabla_modal(), "Propiedades modales y espectrales")
        _pagina_tabla(pdf, res.tabla_fuerzas(), "Fuerzas sísmicas por piso F (t)", 3)
        _pagina_tabla(pdf, res.tabla_cortantes_modales(), "Cortantes por piso y modo V (t)", 3)
        _pagina_tabla(pdf, res.tabla_combinacion(), "Combinación modal y cortante real (t)", 2)

        for fig in figuras.values():
            pdf.savefig(fig, bbox_inches="tight")
    return buffer.getvalue()


def generar_zip_imagenes(res: ResultadosAnalisis, figuras: dict = None) -> bytes:
    """Empaqueta todos los gráficos PNG en un .zip."""
    figuras = figuras or generar_todas_las_figuras(res)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as z:
        for nombre, fig in figuras.items():
            z.writestr(nombre, figura_a_png(fig))
    return buffer.getvalue()
